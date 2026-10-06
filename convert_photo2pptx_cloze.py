import os
import glob
from pptx import Presentation
from pptx.util import Inches
# PIL (Pillow) is used to check the image dimensions before inserting
from PIL import Image 

def bulk_photos_fit_to_ppt(image_folder, output_ppt_name="fitted_images.pptx"):
    # 1. Initialize a widescreen (16:9) presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # Blank layout
    
    # Target maximum dimensions (leaving a 0.5-inch margin on edges)
    max_width = Inches(12.333)
    max_height = Inches(6.5)
    
    # 2. Find all common image formats
    extensions = ('*.jpg', '*.jpeg', '*.png', '*.webp', '*.bmp')
    image_files = []
    for ext in extensions:
        image_files.extend(glob.glob(os.path.join(image_folder, ext)))
        
    image_files.sort()
    
    if not image_files:
        print(f"No images found in folder: {image_folder}")
        return

    print(f"Processing {len(image_files)} photos...")

    for img_path in image_files:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # 3. Calculate aspect ratio using Pillow
        with Image.open(img_path) as img:
            img_width, img_height = img.size
        
        img_aspect = img_width / img_height
        target_aspect = max_width / max_height
        
        # Determine best size to fit without stretching
        if img_aspect > target_aspect:
            # Image is wider than the target box aspect ratio
            fit_width = max_width
            fit_height = max_width / img_aspect
        else:
            # Image is taller than the target box aspect ratio
            fit_height = max_height
            fit_width = max_height * img_aspect
            
        # 4. Center the fitted image perfectly on the 13.333" x 7.5" slide
        left = (prs.slide_width - fit_width) / 2
        top = (prs.slide_height - fit_height) / 2
        
        # 5. Insert the picture with calculated dimensions
        slide.shapes.add_picture(img_path, left, top, width=fit_width, height=fit_height)

    prs.save(output_ppt_name)
    print(f"Successfully created fitted presentation: {output_ppt_name}")

# Example usage:
# bulk_photos_fit_to_ppt(image_folder='my_photos_folder', output_ppt_name='perfect_fit_deck.pptx')

# Example usage:
# Replace 'my_photos_folder' with the path to the folder where your pictures are kept
#bulk_photos_to_ppt(image_folder='/home/liqy/venvs/wronglib/photos/Math', output_ppt_name='math_deck.pptx')
bulk_photos_fit_to_ppt(image_folder='/home/liqy/venvs/wronglib/photos/English/Cloze', output_ppt_name='english_cloze_deck.pptx')


